import base64
import io
import json
import zipfile
from pathlib import Path
from unittest.mock import AsyncMock

import httpx
import pytest
import respx
from docx import Document
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen.canvas import Canvas

from app.config import Settings
from app.parsers import FileParsingError, ParserCancelled, parse_document


def settings(tmp_path: Path, **changes: object) -> Settings:
    return Settings(_env_file=None, data_dir=tmp_path, **changes)


async def test_docx_recursive_tables_headers_stable_source_locations(tmp_path: Path) -> None:
    source = Document()
    paragraph = source.add_paragraph()
    paragraph.add_run("计划投入 ").bold = True
    paragraph.add_run("100 万元。")
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格目标 90%"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "嵌套表格 80%"
    source.sections[0].header.paragraphs[0].text = "年度项目报告"
    path = tmp_path / "source.docx"
    source.save(path)
    first = await parse_document(path, "main", path.name, settings(tmp_path), tmp_path)
    second = await parse_document(path, "main", path.name, settings(tmp_path), tmp_path)
    assert [block.id for block in first.blocks] == [block.id for block in second.blocks]
    assert {block.text for block in first.blocks} == {"计划投入 100 万元。", "表格目标 90%", "嵌套表格 80%", "年度项目报告"}
    main = next(block for block in first.blocks if block.text.startswith("计划"))
    assert main.location["runs"][1]["start"] == len("计划投入 ")
    assert main.location["commentable"] is True
    header = next(block for block in first.blocks if block.text == "年度项目报告")
    assert header.location["commentable"] is False
    assert any(block.kind == "table_cell" and "嵌套" in block.text for block in first.blocks)


async def test_xlsx_formula_cache_missing_is_not_zero_and_chart_values(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "预算"
    sheet.append(["项目", "收入"])
    sheet.append(["甲", 100])
    sheet.append(["乙", "=B2*2"])
    sheet.merge_cells("D1:E1")
    sheet["D1"] = "合并标题"
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=2))
    sheet.add_chart(chart, "G1")
    path = tmp_path / "book.xlsx"
    workbook.save(path)
    parsed = await parse_document(path, "book", path.name, settings(tmp_path), tmp_path)
    assert any(block.text == "100" and block.location["cell"] == "B2" for block in parsed.blocks)
    assert not any(block.location.get("cell") == "B3" for block in parsed.blocks)
    assert any(w.code == "formula_cache_missing" and w.location["cell"] == "B3" for w in parsed.warnings)
    assert any(w.code == "chart_cache_missing" for w in parsed.warnings)
    merged = next(block for block in parsed.blocks if block.text == "合并标题")
    assert merged.location["merged_range"] == "D1:E1"


async def test_pptx_groups_tables_notes_and_native_chart_cache(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "分组里的内容"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "总额"
    table.cell(0, 1).text = "100"
    data = CategoryChartData()
    data.categories = ["一月", "二月"]
    data.add_series("收入", [100, 150])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1), Inches(4), Inches(3), data)
    slide.notes_slide.notes_text_frame.text = "演讲备注 2026"
    path = tmp_path / "slides.pptx"
    presentation.save(path)
    parsed = await parse_document(path, "slides", path.name, settings(tmp_path), tmp_path)
    assert any(block.text == "分组里的内容" and len(block.location["shape_path"]) == 2 for block in parsed.blocks)
    assert any(block.kind == "notes" and "2026" in block.text for block in parsed.blocks)
    assert any(block.kind == "table_cell" and block.text == "100" for block in parsed.blocks)
    assert any(block.kind == "chart" and "150" in block.text and "二月" in block.text for block in parsed.blocks)
    chart = next(block for block in parsed.blocks if block.kind == "chart")
    assert chart.location["slide"] == 1
    assert not any(w.code == "embedded_objects_excluded" for w in parsed.warnings)


async def test_smartart_is_explicitly_reported_as_not_covered(tmp_path: Path) -> None:
    source = Document()
    source.add_paragraph("正文可读取")
    path = tmp_path / "diagram.docx"
    source.save(path)
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("word/diagrams/data1.xml", "<diagram/>")
    parsed = await parse_document(path, "doc", path.name, settings(tmp_path), tmp_path)
    assert any(w.code == "smartart_excluded" and w.incomplete for w in parsed.warnings)


async def test_pdf_text_coordinates_and_scan_coverage_warning(tmp_path: Path) -> None:
    path = tmp_path / "local.pdf"
    canvas = Canvas(str(path))
    canvas.drawString(72, 700, "Budget 100")
    canvas.showPage()
    canvas.showPage()
    canvas.save()
    parsed = await parse_document(path, "pdf", path.name, settings(tmp_path), tmp_path)
    assert any(block.text == "Budget 100" and block.location["page"] == 1 and block.location["bbox"] for block in parsed.blocks)
    assert any(w.code == "pdf_no_text" and w.location["page"] == 2 for w in parsed.warnings)


async def test_ocr_batches_auth_count_validation_and_saved_response(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    for _ in range(12):
        writer.add_blank_page(width=100, height=100)
    writer.write(path)
    counts: list[int] = []

    def respond(request: httpx.Request) -> httpx.Response:
        payload = json.loads(request.content)
        counts.append(len(PdfReader(io.BytesIO(base64.b64decode(payload["file"]))).pages))
        assert payload["fileType"] == 0
        assert payload["restructurePages"] is False
        assert request.headers["Authorization"] == "token secret"
        count = counts[-1] if len(counts) == 1 else 1  # Deliberately omit the last page.
        return httpx.Response(200, json={"errorCode": 0, "result": {"layoutParsingResults": [
            {"markdown": {"text": f"批次 {len(counts)} 第 {i + 1} 页"}, "prunedResult": {}}
            for i in range(count)
        ]}})

    with respx.mock as router:
        router.post("https://ocr.test/layout-parsing").mock(side_effect=respond)
        parsed = await parse_document(path, "ocr", path.name,
                                      settings(tmp_path, paddleocr_url="https://ocr.test", paddleocr_token="secret"), tmp_path)
    assert counts == [10, 2]
    assert [block.location["page"] for block in parsed.blocks] == list(range(1, 12))
    assert any(w.code == "ocr_page_count_mismatch" for w in parsed.warnings)
    assert any(w.code == "ocr_missing_page" and w.location["page"] == 12 for w in parsed.warnings)
    assert len(list(tmp_path.glob("ocr-ocr-*.json"))) == 2


async def test_ocr_failure_does_not_silently_use_local_text(tmp_path: Path) -> None:
    path = tmp_path / "failure.pdf"
    canvas = Canvas(str(path))
    canvas.drawString(72, 700, "Locally extractable")
    canvas.save()
    with respx.mock as router:
        router.post("https://ocr.test/layout-parsing").respond(503, json={"errorMsg": "busy"})
        parsed = await parse_document(path, "pdf", path.name, settings(tmp_path, paddleocr_url="https://ocr.test", llm_max_retries=0), tmp_path)
    assert not parsed.blocks
    assert any(w.code == "ocr_batch_failed" for w in parsed.warnings)


async def test_password_protected_pdf_and_invalid_office_are_errors(tmp_path: Path) -> None:
    path = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("password")
    writer.write(path)
    with pytest.raises(FileParsingError, match="密码"):
        await parse_document(path, "pdf", path.name, settings(tmp_path), tmp_path)
    broken = tmp_path / "bad.docx"
    broken.write_bytes(b"invalid")
    with pytest.raises(FileParsingError):
        await parse_document(broken, "doc", broken.name, settings(tmp_path), tmp_path)


async def test_ocr_retries_transient_errors_but_not_auth(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(path)
    delay = AsyncMock()
    monkeypatch.setattr("app.parsers._retry_delay", delay)
    good = {"result": {"layoutParsingResults": [{"markdown": {"text": "完成"}}]}}
    with respx.mock as router:
        route = router.post("https://ocr.test/layout-parsing").mock(side_effect=[
            httpx.Response(503, json={"errorMsg": "busy"}), httpx.Response(200, json=good)])
        parsed = await parse_document(path, "pdf", path.name, settings(tmp_path, paddleocr_url="https://ocr.test"), tmp_path)
        assert route.call_count == 2
        assert parsed.blocks[0].text == "完成"
    assert delay.await_count == 1
    with respx.mock as router:
        route = router.post("https://ocr.test/layout-parsing").respond(401, json={"errorMsg": "unauthorized"})
        parsed = await parse_document(path, "pdf", path.name, settings(tmp_path, paddleocr_url="https://ocr.test"), tmp_path)
        assert route.call_count == 1
        assert any(w.code == "ocr_batch_failed" for w in parsed.warnings)


async def test_parser_cooperatively_cancels_during_office_and_ocr(tmp_path: Path) -> None:
    document = Document()
    for index in range(10):
        document.add_paragraph(str(index))
    path = tmp_path / "cancel.docx"
    document.save(path)
    checks = 0

    def cancelled() -> bool:
        nonlocal checks
        checks += 1
        return checks > 5

    with pytest.raises(ParserCancelled):
        await parse_document(path, "doc", path.name, settings(tmp_path), tmp_path, cancelled=cancelled)
    pdf_path = tmp_path / "cancel.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(pdf_path)
    done = False

    def respond(request: httpx.Request) -> httpx.Response:
        nonlocal done
        done = True
        return httpx.Response(200, json={"result": {"layoutParsingResults": []}})

    with respx.mock as router:
        router.post("https://ocr.test/layout-parsing").mock(side_effect=respond)
        with pytest.raises(ParserCancelled):
            await parse_document(pdf_path, "pdf", pdf_path.name, settings(tmp_path, paddleocr_url="https://ocr.test"), tmp_path,
                                 cancelled=lambda: done)

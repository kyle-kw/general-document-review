import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from app.config import Settings
from app.exports import export_report
from app.parsers import parse_document


NORMALIZED_REQUIREMENTS = {
    "digest": "normalized-digest", "notes": ["预算按含税口径核对。"],
    "custom_points": [{"id": "custom-1", "title": "核对税费说明",
                       "criterion": "预算应明确包含税费金额及计算口径。", "kind": "required_content"}],
}


def details_row(path: Path) -> dict:
    sheet = load_workbook(path)["问题明细"]
    return {cell.value: sheet.cell(2, cell.column).value for cell in sheet[1]}


def make_issue(block: dict, quote: str, **changes: object) -> dict:
    start = block["text"].index(quote)
    return {
        "id": "issue1", "category": "numeric", "scope": "main", "severity": "high",
        "confidence": "confirmed", "title": "前后金额不一致", "description": "预算应确认",
        "suggestion": "核对预算金额", "evidence": [{"document_id": block["document_id"],
            "block_id": block["id"], "quote": quote, "start": start, "end": start + len(quote)}], **changes,
    }


async def test_docx_exact_comment_anchor_preserves_text_styles_existing_comments(tmp_path: Path) -> None:
    source = Document()
    paragraph = source.add_paragraph()
    first = paragraph.add_run("金额为100万元，")
    first.bold = True
    second = paragraph.add_run("后文为90万元。")
    second.italic = True
    source.add_comment([second], "原有批注", author="原作者")
    source.sections[0].header.paragraphs[0].text = "预算报告"
    path = tmp_path / "main.docx"
    source.save(path)
    parsed = await parse_document(path, "main", path.name, Settings(_env_file=None), tmp_path)
    blocks = [block.model_dump() for block in parsed.blocks]
    paragraph_block = next(block for block in blocks if "金额为" in block["text"])
    header_block = next(block for block in blocks if block["text"] == "预算报告")
    issues = [make_issue(paragraph_block, "100万元"), make_issue(header_block, "预算报告", id="issue2")]
    report = {"review_id": "review1", "status": "completed", "main_document_id": "main",
              "documents": [parsed.model_dump()], "issues": issues}
    artifacts = export_report(report, path, tmp_path / "results")
    assert {item["name"] for item in artifacts} == {"report.json", "report.xlsx", "annotated-main.docx"}
    result = Document(tmp_path / "results" / "annotated-main.docx")
    assert result.paragraphs[0].text == source.paragraphs[0].text
    assert len(result.comments) == 2
    assert result.comments.get(0).text == "原有批注"
    assert any(run.text == "100万元" and run.bold for run in result.paragraphs[0].runs)
    assert any(run.text == "后文为90万元。" and run.italic for run in result.paragraphs[0].runs)
    saved = json.loads((tmp_path / "results" / "report.json").read_text(encoding="utf-8"))
    assert saved["issues"][0]["annotation_status"] == "annotated"
    assert saved["issues"][1]["annotation_status"] == "unavailable"
    workbook = load_workbook(tmp_path / "results" / "report.xlsx")
    assert workbook.sheetnames == ["审查汇总", "问题明细", "解析与覆盖情况"]
    assert workbook["问题明细"]["P2"].value == "annotated"


async def test_cross_document_issue_attaches_to_main_with_both_evidence(tmp_path: Path) -> None:
    source = Document()
    source.add_paragraph("总额100万元")
    path = tmp_path / "main.docx"
    source.save(path)
    parsed = await parse_document(path, "main", path.name, Settings(_env_file=None), tmp_path)
    block = parsed.blocks[0].model_dump()
    issue = make_issue(block, "100万元", scope="cross")
    issue["evidence"].insert(0, {"document_id": "other", "block_id": "other-block", "quote": "90万元"})
    other = {"id": "other", "filename": "other.pdf", "format": "pdf", "warnings": [],
             "blocks": [{"id": "other-block", "document_id": "other", "text": "90万元", "location": {"page": 2}}]}
    report = {"review_id": "cross", "main_document_id": "main", "documents": [parsed.model_dump(), other], "issues": [issue]}
    export_report(report, path, tmp_path / "out")
    result = Document(tmp_path / "out" / "annotated-main.docx")
    comment = result.comments.get(0).text
    assert "other.pdf" in comment and "90万元" in comment and "100万元" in comment
    assert "页 2" in comment


def test_empty_report_and_excel_formula_injection_are_safe(tmp_path: Path) -> None:
    report = {"review_id": "empty", "main_document_id": "main", "documents": [
        {"id": "main", "filename": "=malicious.xlsx", "format": "xlsx", "blocks": [], "warnings": []}], "issues": []}
    export_report(report, tmp_path / "source.xlsx", tmp_path / "empty")
    workbook = load_workbook(tmp_path / "empty" / "report.xlsx")
    assert workbook["问题明细"].max_row == 1
    filename = workbook["审查汇总"]["B4"]
    assert filename.value == "=malicious.xlsx"
    assert filename.data_type == "s"
    assert not (tmp_path / "empty" / "annotated-main.docx").exists()


def test_long_fields_split_without_silent_excel_truncation(tmp_path: Path) -> None:
    block = {"id": "b", "document_id": "main", "text": "证据", "location": {"page": 1}}
    description = "=SUM(A1)" + "长说明" * 15000
    issue = make_issue(block, "证据", description=description)
    report = {"main_document_id": "main", "documents": [{"id": "main", "filename": "a.pdf", "format": "pdf",
        "blocks": [block], "warnings": []}], "issues": [issue]}
    export_report(report, tmp_path / "a.pdf", tmp_path / "out")
    sheet = load_workbook(tmp_path / "out" / "report.xlsx")["问题明细"]
    restored = "".join(str(sheet.cell(row, 7).value or "") for row in range(2, sheet.max_row + 1))
    assert restored == description
    assert sheet["G2"].data_type == "s"


async def test_custom_missing_main_uses_report_only_and_never_anchors_context(tmp_path: Path) -> None:
    source = Document()
    run = source.add_paragraph().add_run("项目预算见附件。")
    run.bold = True
    source.add_comment([run], "原有说明", author="原作者")
    path = tmp_path / "main.docx"
    source.save(path)
    parsed = await parse_document(path, "main", path.name, Settings(_env_file=None), tmp_path)
    block = parsed.blocks[0].model_dump()
    issue = make_issue(block, "项目预算", category="custom", scope="cross", finding_kind="missing",
                       custom_rule_id="custom-1", custom_rule_title="核对税费说明",
                       missing_document_id="main", title="主文件缺少税费说明",
                       absence_coverage=[{"document_id": "main", "reviewed_span_ids": ["span-1"],
                                          "block_ids": [block["id"]], "complete": True, "passes": 2}])
    issue["evidence"][0]["purpose"] = "context"
    issue["evidence"].append({"document_id": "other", "block_id": "other-block", "quote": "税费20万元",
                              "purpose": "quote", "location": {"page": 2}})
    other = {"id": "other", "filename": "预算附件.pdf", "format": "pdf", "warnings": [],
             "blocks": [{"id": "other-block", "document_id": "other", "text": "税费20万元", "location": {"page": 2}}]}
    report = {"review_id": "custom-missing", "main_document_id": "main", "status": "completed",
              "documents": [parsed.model_dump(), other], "issues": [issue],
              "review_requirements": NORMALIZED_REQUIREMENTS, "requirements_policy_version": "test-policy"}
    artifacts = export_report(report, path, tmp_path / "out")
    assert {item["name"] for item in artifacts} == {"report.json", "report.xlsx", "annotated-main.docx"}
    saved = json.loads((tmp_path / "out" / "report.json").read_text(encoding="utf-8"))
    saved_issue = saved["issues"][0]
    assert saved_issue["finding_kind"] == "missing"
    assert saved_issue["annotation_status"] == "unavailable"
    assert "真实主文件问题引文" in saved_issue["annotation_reason"]
    assert saved["review_requirements"] == NORMALIZED_REQUIREMENTS
    assert saved_issue["evidence"] == issue["evidence"]
    annotated = Document(tmp_path / "out" / "annotated-main.docx")
    assert annotated.paragraphs[0].text == source.paragraphs[0].text
    assert annotated.paragraphs[0].runs[0].bold
    assert len(annotated.comments) == 1
    assert annotated.comments.get(0).text == "原有说明"
    row = details_row(tmp_path / "out" / "report.xlsx")
    assert row["问题 ID"] == saved_issue["id"]
    assert row["类型"] == "专项审查"
    assert row["规则 ID"] == "custom-1"
    assert row["审查要点"] == "核对税费说明"
    assert row["问题性质"] == "内容缺失"
    assert row["主文件/证据一"] == "main.docx"
    assert row["位置一"] is None and row["引文一"] is None
    assert row["关联文件/证据二"] == "预算附件.pdf"
    assert row["引文二"] == "税费20万元"
    assert row["缺失文件"] == "main.docx"
    assert "已核对全部可提取内容" in row["缺失内容核对覆盖"]
    assert "2 轮核查" in row["缺失内容核对覆盖"]
    assert row["批注状态"] == saved_issue["annotation_status"]
    assert json.loads(row["其他证据"])[0]["purpose"] == "context"
    workbook = load_workbook(tmp_path / "out" / "report.xlsx")
    summary = {row[0].value: row[1].value for row in workbook["审查汇总"]}
    assert summary["审查要点 1：核对税费说明"] == NORMALIZED_REQUIREMENTS["custom_points"][0]["criterion"]
    assert summary["审查备注 1"] == "预算按含税口径核对。"
    assert summary["审查要求策略版本"] == "test-policy"
    coverage = list(workbook["解析与覆盖情况"].values)
    assert any(row[3] == "缺失核对完整" and row[4] == "custom_absence_coverage" and "2 轮核查" in row[6] for row in coverage[1:])


async def test_custom_cross_missing_attachment_comments_real_main_quote_with_absence_coverage(tmp_path: Path) -> None:
    source = Document()
    run = source.add_paragraph().add_run("项目税费20万元。")
    run.italic = True
    path = tmp_path / "main.docx"
    source.save(path)
    parsed = await parse_document(path, "main", path.name, Settings(_env_file=None), tmp_path)
    block = parsed.blocks[0].model_dump()
    issue = make_issue(block, "税费20万元", category="custom", scope="cross", finding_kind="missing",
                       custom_rule_id="custom-1", custom_rule_title="核对税费说明", missing_document_id="other",
                       title="附件缺少主文件列明的税费", description="主文件列明税费，附件未找到对应金额。",
                       absence_coverage=[{"document_id": "other", "reviewed_span_ids": ["span-other"],
                                          "block_ids": ["other-block"], "complete": True, "passes": 2}])
    issue["evidence"].append({"document_id": "main", "block_id": block["id"], "quote": "项目",
                              "purpose": "quote", "location": block["location"]})
    other = {"id": "other", "filename": "成本.xlsx", "format": "xlsx", "warnings": [],
             "blocks": [{"id": "other-block", "document_id": "other", "text": "建设支出100万元",
                         "location": {"sheet": "成本", "cell": "A1"}}]}
    report = {"main_document_id": "main", "documents": [parsed.model_dump(), other], "issues": [issue],
              "review_requirements": NORMALIZED_REQUIREMENTS}
    export_report(report, path, tmp_path / "out")
    saved_issue = json.loads((tmp_path / "out" / "report.json").read_text(encoding="utf-8"))["issues"][0]
    annotated = Document(tmp_path / "out" / "annotated-main.docx")
    comment = annotated.comments.get(0).text
    assert saved_issue["annotation_status"] == "annotated"
    assert "审查要点：核对税费说明" in comment and "问题性质：内容缺失" in comment
    assert "税费20万元" in comment
    assert "成本.xlsx：已核对全部可提取内容" in comment
    assert "未找到所需内容" in comment
    assert "2 轮核查" in comment
    assert "建设支出100万元" not in comment
    assert "单元格 A1" not in comment
    assert annotated.paragraphs[0].text == source.paragraphs[0].text
    assert any(run.text == "税费20万元" and run.italic for run in annotated.paragraphs[0].runs)
    row = details_row(tmp_path / "out" / "report.xlsx")
    assert row["问题 ID"] == saved_issue["id"]
    assert row["主文件/证据一"] == "main.docx" and row["引文一"] == "税费20万元"
    assert row["关联文件/证据二"] == "成本.xlsx"
    assert row["位置二"] is None and row["引文二"] is None
    assert json.loads(row["其他证据"])[0]["quote"] == "项目"
    assert row["缺失文件"] == "成本.xlsx"
    assert row["批注状态"] == saved_issue["annotation_status"]


async def test_custom_missing_main_without_any_main_evidence_remains_unannotated(tmp_path: Path) -> None:
    source = Document()
    source.add_paragraph("项目方案概述")
    path = tmp_path / "main.docx"
    source.save(path)
    parsed = await parse_document(path, "main", path.name, Settings(_env_file=None), tmp_path)
    other_block = {"id": "other-block", "document_id": "other", "text": "税费20万元", "location": {"page": 1}}
    issue = make_issue(other_block, "税费20万元", category="custom", scope="cross", finding_kind="missing",
                       custom_rule_id="custom-1", custom_rule_title="核对税费说明", missing_document_id="main",
                       absence_coverage=[{"document_id": "main", "reviewed_span_ids": ["span-main"],
                                          "block_ids": [parsed.blocks[0].id], "complete": False}])
    report = {"main_document_id": "main", "documents": [parsed.model_dump(),
              {"id": "other", "filename": "附件.pdf", "format": "pdf", "blocks": [other_block], "warnings": []}],
              "issues": [issue]}
    export_report(report, path, tmp_path / "out")
    saved_issue = json.loads((tmp_path / "out" / "report.json").read_text(encoding="utf-8"))["issues"][0]
    assert saved_issue["annotation_status"] == "unavailable"
    assert "没有真实主文件问题引文" in saved_issue["annotation_reason"]
    assert len(Document(tmp_path / "out" / "annotated-main.docx").comments) == 0
    row = details_row(tmp_path / "out" / "report.xlsx")
    assert row["位置一"] is None and row["引文一"] is None
    assert "审查覆盖不完整，需人工确认" in row["缺失内容核对覆盖"]
    summary = {row[0].value: row[1].value for row in load_workbook(tmp_path / "out" / "report.xlsx")["审查汇总"]}
    assert summary["审查覆盖"] == "不完整，详见解析与覆盖情况"


async def test_custom_violation_records_rule_and_anchors_only_the_actual_quote(tmp_path: Path) -> None:
    source = Document()
    source.add_paragraph("项目预算为未税金额100万元。")
    path = tmp_path / "main.docx"
    source.save(path)
    parsed = await parse_document(path, "main", path.name, Settings(_env_file=None), tmp_path)
    issue = make_issue(parsed.blocks[0].model_dump(), "未税金额100万元", category="custom", finding_kind="violation",
                       custom_rule_id="custom-1", custom_rule_title="核对税费说明")
    report = {"main_document_id": "main", "documents": [parsed.model_dump()], "issues": [issue],
              "review_requirements": NORMALIZED_REQUIREMENTS}
    export_report(report, path, tmp_path / "out")
    saved_issue = json.loads((tmp_path / "out" / "report.json").read_text(encoding="utf-8"))["issues"][0]
    comment = Document(tmp_path / "out" / "annotated-main.docx").comments.get(0).text
    assert "专项审查" in comment and "审查要点：核对税费说明" in comment
    assert "问题性质：内容不符" in comment and "未税金额100万元" in comment
    row = details_row(tmp_path / "out" / "report.xlsx")
    assert row["审查要点"] == saved_issue["custom_rule_title"]
    assert row["问题性质"] == "内容不符"
    assert row["缺失文件"] is None and row["缺失内容核对覆盖"] is None
    assert row["批注状态"] == saved_issue["annotation_status"] == "annotated"


async def test_zero_issues_keeps_normalized_requirements_and_docx_output(tmp_path: Path) -> None:
    source = Document()
    source.add_paragraph("项目说明。")
    path = tmp_path / "main.docx"
    source.save(path)
    parsed = await parse_document(path, "main", path.name, Settings(_env_file=None), tmp_path)
    report = {"main_document_id": "main", "documents": [parsed.model_dump()], "issues": [],
              "review_requirements": NORMALIZED_REQUIREMENTS}
    artifacts = export_report(report, path, tmp_path / "out")
    assert len(artifacts) == 3
    saved = json.loads((tmp_path / "out" / "report.json").read_text(encoding="utf-8"))
    assert saved["issues"] == [] and saved["review_requirements"] == NORMALIZED_REQUIREMENTS
    document = Document(tmp_path / "out" / "annotated-main.docx")
    assert document.paragraphs[0].text == "项目说明。" and len(document.comments) == 0
    workbook = load_workbook(tmp_path / "out" / "report.xlsx")
    assert workbook["问题明细"].max_row == 1
    assert workbook.sheetnames == ["审查汇总", "问题明细", "解析与覆盖情况"]

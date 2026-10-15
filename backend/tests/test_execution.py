"""Real extraction/orchestration/export/storage integration with a deterministic model double."""
import json
import os
from pathlib import Path
import re
from typing import Any
from uuid import uuid4

from docx import Document
import fakeredis
from openpyxl import Workbook, load_workbook
import pytest
from rq import Queue, SimpleWorker
from rq.serializers import JSONSerializer
from rq.timeouts import TimerDeathPenalty
from redis import Redis

from app import codex_client, jobs
from app.codex_client import AgentError, ReviewCancelled
from app.config import Settings
from app.storage import review_directory
from app.store import Store


def source_quote(span: dict[str, Any], quote: str, local_start: int | None = None) -> dict[str, Any]:
    position = span["text"].index(quote) if local_start is None else local_start
    start = span["offset"] + position
    return {"document_id": span["document_id"], "block_id": span["block_id"], "quote": quote,
            "start": start, "end": start + len(quote)}


def fact_quote(fact: dict[str, Any]) -> dict[str, Any]:
    evidence = fact["evidence"][0]
    return {key: evidence[key] for key in ("document_id", "block_id", "quote", "start", "end")}


def proposed(category: str, evidence: list[dict[str, Any]], title: str) -> dict[str, Any]:
    return {"category": category, "severity": "medium", "confidence": "confirmed", "title": title,
            "description": title, "suggestion": "核对并统一表述", "evidence": evidence}


class DeterministicAgents:
    """Return real source quotes, never bypassing pipeline validation or exports."""
    calls: list[tuple[str, dict[str, Any]]] = []
    fail = False
    cancel: Any = None

    def __init__(self, settings: Settings, work_dir: Path, cancelled: Any, *, source_files=None) -> None:
        self.context_chars = settings.llm_context_chars
        self.concurrency = settings.agent_concurrency
        self.cache_identity = "execution-regression-double-v1"

    async def __aenter__(self):
        return self

    async def __aexit__(self, *args):
        return None

    async def run_agent(self, instructions: str, payload: dict[str, Any], schema: dict) -> dict:
        role = next(role for role in ("proofreader", "facts", "numeric", "semantic", "cross", "verifier")
                    if f"[{role}]" in instructions)
        type(self).calls.append((role, payload))
        if type(self).cancel is not None:
            type(self).cancel()
            raise ReviewCancelled("模拟模型 turn 响应用户取消")
        if type(self).fail:
            raise AgentError("模拟网关不可用")
        if role == "proofreader":
            issues = [proposed("typo", [source_quote(span, "计化")], "计化应为计划")
                      for span in payload["spans"] if "计化" in span["text"]]
            return {"reviewed_span_ids": payload["target_span_ids"], "issues": issues}
        if role == "facts":
            facts = []
            for span in payload["spans"]:
                for match in re.finditer(r"\d+", span["text"]):
                    facts.append({"entity": "项目", "metric": "总预算", "value": match.group(), "unit": "万元",
                                  "period": "", "scope": "", "claim": f"项目总预算为{match.group()}万元",
                                  "evidence": [source_quote(span, match.group(), match.start())]})
            return {"reviewed_span_ids": payload["target_span_ids"], "facts": facts}
        if role == "verifier":
            return {"decisions": [{"candidate_id": item["candidate_id"], "keep": True,
                                   "confidence": "confirmed", "reason": "双方证据匹配原文"}
                                  for item in payload["candidates"]]}
        if role == "numeric":
            facts = payload["left_facts"] + payload["right_facts"]
            first = next((fact for fact in facts if fact["value"] == "100"), None)
            second = next((fact for fact in facts if fact["value"] == "120"), None)
            return {"issues": [proposed("numeric", [fact_quote(first), fact_quote(second)], "主文件预算不一致")]
                    if first and second else []}
        if role == "cross":
            first = next((fact for fact in payload["main_facts"] if fact["value"] == "100"), None)
            second = next((fact for fact in payload["other_facts"] if fact["value"] == "90"), None)
            return {"issues": [proposed("numeric", [fact_quote(first), fact_quote(second)], "附件与主文件预算不一致")]
                    if first and second else []}
        return {"issues": []}


@pytest.fixture
def setup_review(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    settings = Settings(_env_file=None, data_dir=tmp_path / "data", static_dir=tmp_path / "dist")
    store = Store(settings.database_path)
    store.initialize()
    owner = store.add_owner("browser-token-" + "x" * 32)
    DeterministicAgents.calls = []
    DeterministicAgents.fail = False
    DeterministicAgents.cancel = None
    monkeypatch.setattr(codex_client, "CodexAgentClient", DeterministicAgents)

    def create(*, reference: str = "valid", invalid_main: bool = False) -> tuple[str, Path]:
        identifier = uuid4().hex
        directory = review_directory(settings.data_dir, identifier)
        originals = directory / "originals"
        originals.mkdir(parents=True)
        main = originals / "main.docx"
        if invalid_main:
            main.write_bytes(b"broken DOCX")
        else:
            document = Document()
            document.add_paragraph("项目计化。")
            document.add_paragraph("总预算为100万元。")
            document.add_paragraph("总预算为120万元。")
            document.save(main)
        files = [{"id": "main", "filename": "主文件.docx", "path": "originals/main.docx", "is_main": True}]
        if reference != "none":
            path = originals / "reference.xlsx"
            if reference == "valid":
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(["总预算（万元）", "其他内容"])
                sheet.append([90, "不应校对附表错别子"])
                workbook.save(path)
            elif reference == "corrupt":
                path.write_bytes(b"broken XLSX")
            # 'missing' intentionally records an absent source to verify per-attachment failure isolation.
            files.append({"id": "reference", "filename": "参考文件.xlsx", "path": "originals/reference.xlsx", "is_main": False})
        assert store.create(identifier, owner, uuid4().hex, files)
        return identifier, directory

    return settings, store, create


async def test_execute_review_real_docx_and_xlsx_pipeline_artifacts_and_durable_state(setup_review) -> None:
    settings, store, create = setup_review
    identifier, directory = create()
    await jobs.execute_review(identifier, settings, store)
    item = store.get(identifier, private=True)
    assert item["status"] == "completed", item.get("error")
    assert item["issue_count"] == 3
    assert item["finished_at"] and item["expires_at"]
    assert {entry["name"] for entry in item["artifacts"]} == {"report.json", "report.xlsx", "annotated-main.docx"}
    assert not list(directory.glob("output-*.tmp"))
    artifacts = {entry["kind"]: directory / entry["path"] for entry in item["artifacts"]}
    report = json.loads(artifacts["json"].read_text(encoding="utf-8"))
    assert report["status"] == item["status"]
    saved_issues = store.issues(identifier)["items"]
    assert saved_issues == report["issues"]
    assert all(issue["annotation_status"] == "annotated" for issue in saved_issues)
    assert {issue["category"] for issue in saved_issues} == {"typo", "numeric"}
    cross = next(issue for issue in saved_issues if issue["scope"] == "cross")
    assert {evidence["document_id"] for evidence in cross["evidence"]} == {"main", "reference"}
    assert cross["confidence"] == "needs_confirmation"
    assert "基准有歧义" in cross["description"]
    annotated = Document(artifacts["annotated"])
    assert len(annotated.comments) == 3
    assert [paragraph.text for paragraph in annotated.paragraphs] == ["项目计化。", "总预算为100万元。", "总预算为120万元。"]
    spreadsheet = load_workbook(artifacts["excel"])
    assert spreadsheet["问题明细"].max_row == 4
    assert spreadsheet["解析与覆盖情况"].max_row >= 3
    proofreading = [payload for role, payload in DeterministicAgents.calls if role == "proofreader"]
    assert proofreading and all(payload["document_id"] == "main" for payload in proofreading)
    roles = [role for role, _ in DeterministicAgents.calls]
    assert roles.index("verifier") < roles.index("cross")
    assert (directory / "work" / "checkpoints" / "main-fact-index.json").exists()
    reopened = Store(settings.database_path)
    assert reopened.get(identifier)["issue_count"] == 3
    assert "path" not in reopened.get(identifier)["artifacts"][0]


async def test_primary_parse_failure_is_terminal_and_never_calls_model(setup_review) -> None:
    settings, store, create = setup_review
    identifier, directory = create(invalid_main=True)
    await jobs.execute_review(identifier, settings, store)
    item = store.get(identifier)
    assert item["status"] == "failed"
    assert "主文件解析失败" in item["error"]
    assert not item["artifacts"]
    assert not DeterministicAgents.calls
    assert not list(directory.glob("output-*"))


@pytest.mark.parametrize("reference", ["corrupt", "missing"])
async def test_attachment_failure_keeps_main_report_partial(setup_review, reference: str) -> None:
    settings, store, create = setup_review
    identifier, directory = create(reference=reference)
    await jobs.execute_review(identifier, settings, store)
    item = store.get(identifier, private=True)
    assert item["status"] == "partial", item.get("error")
    assert item["issue_count"] == 2
    assert any(warning["code"] == "parse_failed" and warning["document_id"] == "reference" for warning in item["warnings"])
    report_path = directory / next(entry["path"] for entry in item["artifacts"] if entry["kind"] == "json")
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["status"] == "partial"
    assert all(issue["scope"] == "main" for issue in report["issues"])


@pytest.mark.parametrize("when", ["queued", "running"])
async def test_cancellation_leaves_no_published_output(setup_review, when: str) -> None:
    settings, store, create = setup_review
    identifier, directory = create()
    if when == "queued":
        store.update(identifier, cancel_requested=1)
    else:
        DeterministicAgents.cancel = lambda: store.update(identifier, cancel_requested=1)
    await jobs.execute_review(identifier, settings, store)
    item = store.get(identifier)
    assert item["status"] == "cancelled"
    assert item["expires_at"] and not item["artifacts"]
    assert store.issues(identifier)["total"] == 0
    assert not list(directory.glob("output-*"))
    if when == "queued":
        assert not DeterministicAgents.calls


async def test_model_failure_is_terminal_without_half_published_results(setup_review) -> None:
    settings, store, create = setup_review
    identifier, directory = create()
    DeterministicAgents.fail = True
    await jobs.execute_review(identifier, settings, store)
    item = store.get(identifier)
    assert item["status"] == "failed"
    assert "模型审查失败" in item["error"]
    assert not item["artifacts"] and not list(directory.glob("output-*"))


async def test_cancellation_between_snapshot_and_claim_is_immediately_terminal(setup_review, monkeypatch) -> None:
    settings, store, create = setup_review
    identifier, directory = create()
    original_get = store.get
    captured = False

    def cancel_after_snapshot(*args, **kwargs):
        nonlocal captured
        item = original_get(*args, **kwargs)
        if not captured:
            captured = True
            store.update(identifier, cancel_requested=1)
        return item

    monkeypatch.setattr(store, "get", cancel_after_snapshot)
    await jobs.execute_review(identifier, settings, store)
    item = original_get(identifier)
    assert item["status"] == "cancelled"
    assert item["finished_at"] and item["expires_at"]
    assert not DeterministicAgents.calls and not item["artifacts"]
    assert not list(directory.glob("output-*"))


async def test_retry_does_not_reuse_older_parser_cache(setup_review) -> None:
    settings, store, create = setup_review
    identifier, directory = create(reference="none")
    workspace = directory / "work"
    workspace.mkdir()
    import hashlib
    from app.models import Block, ParsedDocument
    source = directory / "originals" / "main.docx"
    fingerprint = hashlib.sha256(source.read_bytes()).hexdigest()
    old_key = hashlib.sha256(f"parser-v1:{fingerprint}:{settings.paddleocr_url}:{settings.paddleocr_batch_pages}".encode()).hexdigest()
    stale = ParsedDocument(id="main", filename="主文件.docx", format="docx", blocks=[
        Block(id="stale", document_id="main", text="Stale incomplete extraction", kind="paragraph", location={})])
    jobs.atomic_json(workspace / "parsed-main.json", {"key": old_key, "document": stale.model_dump()})
    await jobs.execute_review(identifier, settings, store)
    item = store.get(identifier, private=True)
    assert item["status"] == "completed", item["error"]
    assert item["issue_count"] == 2
    assert "项目计化。" in [block["text"] for block in item["documents"][0]["blocks"]]
    assert "stale" not in [block["id"] for block in item["documents"][0]["blocks"]]


async def test_retry_reuses_complete_main_checkpoints_and_replaces_registered_artifacts(setup_review) -> None:
    settings, store, create = setup_review
    identifier, directory = create(reference="corrupt")
    await jobs.execute_review(identifier, settings, store)
    original = store.get(identifier, private=True)
    assert original["status"] == "partial" and original["issue_count"] == 2
    parsed_main = directory / "work" / "parsed-main.json"
    original_mtime = parsed_main.stat().st_mtime_ns
    workbook = Workbook()
    workbook.active.append(["总预算（万元）"])
    workbook.active.append([90])
    workbook.save(directory / "originals" / "reference.xlsx")
    store.update(identifier, status="queued", attempt=2, rq_job_id=f"{identifier}-2", cancel_requested=0,
                 finished_at=None, expires_at=None)
    DeterministicAgents.calls = []
    await jobs.execute_review(identifier, settings, store)
    retried = store.get(identifier, private=True)
    assert retried["status"] == "completed" and retried["issue_count"] == 3
    assert parsed_main.stat().st_mtime_ns == original_mtime
    assert not any(role == "proofreader" for role, _ in DeterministicAgents.calls)
    assert all(payload["document_id"] == "reference" for role, payload in DeterministicAgents.calls if role == "facts")
    assert {artifact["path"] for artifact in retried["artifacts"]}.isdisjoint(
        artifact["path"] for artifact in original["artifacts"])
    assert all((directory / artifact["path"]).exists() for artifact in retried["artifacts"])


@pytest.mark.parametrize("redis_backend", ["fake", "real"])
def test_rq_json_job_runs_actual_execution_entry_point(setup_review, monkeypatch: pytest.MonkeyPatch, redis_backend: str) -> None:
    settings, store, create = setup_review
    identifier, directory = create()
    monkeypatch.setattr(jobs, "Settings", lambda: settings)
    if redis_backend == "real":
        url = os.environ.get("TEST_REDIS_URL")
        if not url:
            pytest.skip("Set TEST_REDIS_URL to an isolated test Redis for real queue validation")
        connection = Redis.from_url(url)
    else:
        connection = fakeredis.FakeRedis()
    queue = Queue("execution-test-" + identifier, connection=connection, serializer=JSONSerializer)
    job = queue.enqueue("app.jobs.run_review_job", identifier, job_id=f"{identifier}-1", job_timeout=30)

    class TimerWorker(SimpleWorker):
        death_penalty_class = TimerDeathPenalty

    worker = TimerWorker([queue], connection=connection, serializer=JSONSerializer)
    worker.work(burst=True, with_scheduler=False, logging_level="WARNING")
    job.refresh()
    assert job.is_finished, job.exc_info
    assert job.return_value() == {"id": identifier, "status": "completed"}
    assert store.get(identifier)["issue_count"] == 3
    assert len(store.get(identifier)["artifacts"]) == 3
    job.delete()
    queue.delete()
    connection.close()

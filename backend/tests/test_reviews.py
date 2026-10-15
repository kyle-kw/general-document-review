from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

from docx import Document
import fakeredis
from fastapi.testclient import TestClient
import pytest
from rq import Queue
from rq.serializers import JSONSerializer

from app.config import Settings
from app.job_queue import get_queue
from app.main import create_app
from app.maintenance import maintain
from app.storage import review_directory


@pytest.fixture
def application(tmp_path: Path, monkeypatch):
    settings = Settings(_env_file=None, data_dir=tmp_path / "data", static_dir=tmp_path / "dist")
    app = create_app(settings)
    connection = fakeredis.FakeRedis()
    monkeypatch.setattr("app.maintenance.make_connection", lambda *a, **k: connection)
    queue = Queue(settings.rq_queue, connection=connection, serializer=JSONSerializer)
    app.dependency_overrides[get_queue] = lambda: queue
    app.state.test_queue = queue
    return app


def document_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("项目预算为100万元。项目预算为120万元。")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def submit(client: TestClient, key: str = "test"):
    return client.post("/api/v1/reviews", files={"main_file": ("预算.docx", document_bytes())},
                       headers={"Idempotency-Key": key})


def test_upload_is_durable_idempotent_and_queued_as_json(application):
    with TestClient(application) as client:
        assert submit(client).status_code == 401
        response = client.post("/api/v1/session")
        assert "HttpOnly" in response.headers["set-cookie"]
        assert "SameSite=strict" in response.headers["set-cookie"]
        first = submit(client)
        assert first.status_code == 202, first.text
        review = first.json()
        assert review["status"] == "queued"
        assert "path" not in review["files"][0]
        assert "owner_id" not in review
        assert review["cancel_requested"] is False
        assert submit(client).json()["id"] == review["id"]
        assert application.state.test_queue.count == 1
        job = application.state.test_queue.fetch_job(f"{review['id']}-1")
        assert job.func_name == "app.jobs.run_review_job" and list(job.args) == [review["id"]]
        assert client.get("/api/v1/reviews").json()["total"] == 1
        downloaded = client.get(f"/api/v1/reviews/{review['id']}/files/{review['files'][0]['id']}").content
        assert "项目预算为100万元" in Document(BytesIO(downloaded)).paragraphs[0].text


def test_browser_isolation_covers_every_read_and_mutation(application):
    with TestClient(application) as first, TestClient(application) as second:
        first.post("/api/v1/session")
        second.post("/api/v1/session")
        review = submit(first).json()
        identifier = review["id"]
        assert second.get("/api/v1/reviews").json()["items"] == []
        for suffix in ("", "/issues", f"/files/{review['files'][0]['id']}", "/artifacts/anything"):
            assert second.get(f"/api/v1/reviews/{identifier}{suffix}").status_code == 404
        for action in ("cancel", "retry"):
            assert second.post(f"/api/v1/reviews/{identifier}/{action}").status_code == 404
        assert second.delete(f"/api/v1/reviews/{identifier}").status_code == 404
        assert first.delete(f"/api/v1/reviews/{identifier}").status_code == 409
        assert first.post(f"/api/v1/reviews/{identifier}/cancel").json()["status"] == "cancelled"
        assert first.delete(f"/api/v1/reviews/{identifier}").status_code == 204
        assert first.get(f"/api/v1/reviews/{identifier}").status_code == 404


def test_validation_origin_and_missing_static(application):
    with TestClient(application) as client:
        assert client.post("/api/v1/session", headers={"Origin": "https://attacker.example"}).status_code == 403
        client.post("/api/v1/session")
        assert client.post("/api/v1/reviews", files={"main_file": ("fake.docx", b"bad")}).status_code == 400
        assert client.post("/api/v1/reviews", files={"main_file": ("old.ppt", b"bad")}).status_code == 415
        assert client.get("/data/app.sqlite3").status_code == 404
        assert client.get("/api/v1/anything").status_code == 404
        assert client.get("/assets/missing.js").status_code == 404
        assert client.get("/api/v1/reviews").json()["total"] == 0


def test_history_survives_redis_expiry_and_retry_preserves_sources(application):
    with TestClient(application) as client:
        client.post("/api/v1/session")
        review = submit(client).json()
        identifier = review["id"]
        store = application.state.store
        store.finish(identifier, "failed", 30, error="测试中断")
        application.state.test_queue.fetch_job(f"{identifier}-1").delete()
        assert client.get(f"/api/v1/reviews/{identifier}").json()["status"] == "failed"
        assert client.post(f"/api/v1/reviews/{identifier}/retry").status_code == 202
        assert client.get(f"/api/v1/reviews/{identifier}").json()["attempt"] == 2
        assert client.post(f"/api/v1/reviews/{identifier}/retry").status_code == 409


def test_retention_only_cleans_expired_terminal_records(application, monkeypatch):
    import app.maintenance as maintenance
    with TestClient(application) as client:
        client.post("/api/v1/session")
        expired = submit(client, "expired").json()["id"]
        active = submit(client, "active").json()["id"]
        store = application.state.store
        store.finish(expired, "completed", 30)
        yesterday = (datetime.now(timezone.utc) - timedelta(days=1)).isoformat()
        store.update(expired, expires_at=yesterday)
        store.update(active, expires_at=yesterday)
        monkeypatch.setattr(maintenance, "make_connection", lambda _: application.state.test_queue.connection)
        maintain(application.state.settings, store)
        assert store.get(expired) is None
        assert not review_directory(application.state.settings.data_dir, expired).exists()
        assert store.get(active)["status"] == "queued"


def test_queue_outage_preserves_retryable_history_without_queued_state(application, monkeypatch):
    from redis.exceptions import ConnectionError
    with TestClient(application) as client:
        client.post("/api/v1/session")
        monkeypatch.setattr(application.state.test_queue, "enqueue", lambda *a, **k: (_ for _ in ()).throw(ConnectionError()))
        assert submit(client).status_code == 503
        history = client.get("/api/v1/reviews").json()
        assert history["total"] == 1 and history["items"][0]["status"] == "failed"
        file = history["items"][0]["files"][0]
        assert client.get(f"/api/v1/reviews/{history['items'][0]['id']}/files/{file['id']}").status_code == 200


def test_stale_maintenance_snapshot_cannot_delete_or_finish_retried_task(application):
    with TestClient(application) as client:
        client.post("/api/v1/session")
        identifier = submit(client).json()["id"]
        store = application.state.store
        store.finish(identifier, "failed", 30)
        store.update(identifier, expires_at=(datetime.now(timezone.utc) - timedelta(days=1)).isoformat())
        assert identifier in store.expired()  # Maintenance took its snapshot.
        assert client.post(f"/api/v1/reviews/{identifier}/retry").status_code == 202
        removed = []
        assert not store.delete_terminal(identifier, lambda: removed.append(True), expired_only=True)
        assert not store.finish_current(identifier, f"{identifier}-1", 30)
        assert removed == [] and store.get(identifier)["status"] == "queued"


def test_cancelling_queued_snapshot_does_not_publish_terminal_after_worker_claim(application, monkeypatch):
    with TestClient(application) as client:
        client.post("/api/v1/session")
        identifier = submit(client).json()["id"]
        store = application.state.store
        job = application.state.test_queue.fetch_job(f"{identifier}-1")
        original_cancel = job.cancel

        def claim_then_cancel():
            store.update(identifier, status="running")
            original_cancel()

        monkeypatch.setattr(job, "cancel", claim_then_cancel)
        monkeypatch.setattr(application.state.test_queue, "fetch_job", lambda _: job)
        assert client.post(f"/api/v1/reviews/{identifier}/cancel").json()["status"] == "running"
        assert store.get(identifier, private=True)["cancel_requested"] == 1
        assert client.get(f"/api/v1/reviews/{identifier}").json()["cancel_requested"] is True
        assert client.delete(f"/api/v1/reviews/{identifier}").status_code == 409


def test_lost_enqueue_response_never_deletes_running_sources_or_overwrites_retry(application, monkeypatch):
    from redis.exceptions import ConnectionError
    with TestClient(application) as client:
        client.post("/api/v1/session")
        queue = application.state.test_queue
        store = application.state.store
        original_enqueue = queue.enqueue

        def commit_claim_lose_response(*args, **kwargs):
            original_enqueue(*args, **kwargs)
            store.update(args[1], status="running")
            raise ConnectionError("simulated lost response after accepted enqueue")

        monkeypatch.setattr(queue, "enqueue", commit_claim_lose_response)
        response = submit(client)
        assert response.status_code == 202 and response.json()["status"] == "running"
        review = response.json()
        path = f"/api/v1/reviews/{review['id']}/files/{review['files'][0]['id']}"
        assert client.get(path).status_code == 200
        store.finish(review["id"], "failed", 30)
        response = client.post(f"/api/v1/reviews/{review['id']}/retry")
        assert response.status_code == 202 and response.json()["status"] == "running"
        assert client.get(path).status_code == 200

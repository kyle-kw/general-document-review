"""Pinned OfficeCLI/skill integration, including actual Linux namespace execution."""

import asyncio
import hashlib
import json
from pathlib import Path
import sys
import zipfile

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
import pytest

from app.config import Settings
from app.sandbox import (
    OFFICECLI_SHA256, OFFICECLI_SKILL, OFFICECLI_VERSION, WorkspaceSandbox,
    create_workspace, officecli_skill_identity, stage_officecli_skill, validate_operation,
)


@pytest.mark.parametrize("argv", [
    ["officecli"], ["officecli", "install"], ["officecli", "update"],
    ["officecli", "config", "autoUpdate", "true"], ["officecli", "plugins", "install", "pdf"],
    ["officecli", "mcp"], ["officecli", "skills", "install"], ["officecli", "open", "a.docx"],
    ["officecli", "watch", "a.docx"], ["officecli", "__resident-serve__"],
    ["officecli", "view", "a.docx", "html"], ["officecli", "view", "a.docx", "screenshot"],
    ["officecli", "view", "a.docx", "text", "--browser"],
])
def test_officecli_excludes_install_plugins_background_services_and_browser(argv):
    with pytest.raises(ValueError, match="OfficeCLI"):
        validate_operation("run_command", {"argv": argv})


def test_officecli_accepts_argument_arrays_for_inspection_and_scratch_experiments():
    for argv in (["officecli", "--version"], ["officecli", "help", "docx", "paragraph", "--json"],
                 ["officecli", "get", "/workspace/input/a.xlsx", "/Sheet1/A1", "--json"],
                 ["officecli", "view", "/workspace/input/a.pptx", "text", "--json"],
                 ["officecli", "create", "/workspace/scratch/test.docx"]):
        validate_operation("run_command", {"argv": argv})


def test_skill_is_trusted_review_adaptation_and_versioned_with_checkpoints(tmp_path):
    workspace = create_workspace(tmp_path, {})
    metadata = stage_officecli_skill(workspace)
    content = (workspace / metadata["path"]).read_bytes()
    assert content == OFFICECLI_SKILL.read_bytes()
    assert metadata["version"] == OFFICECLI_VERSION == "1.0.157"
    assert metadata["sha256"] == hashlib.sha256(content).hexdigest()
    assert metadata["identity"] == officecli_skill_identity()
    assert not (workspace / "input/skills/officecli/upstream").exists()
    for name in ("docx.md", "xlsx.md", "pptx.md"):
        assert (workspace / "input/skills/officecli/references" / name).read_bytes() == (OFFICECLI_SKILL.parent / "upstream" / name).read_bytes()
    text = content.decode("utf-8")
    assert "computedValue" in text and "缺失缓存" in text and "不可信数据" in text
    assert "curl -" not in text and "officecli install" not in text


def test_vendored_official_skills_and_notices_match_fixed_commit():
    root = OFFICECLI_SKILL.parent
    provenance = json.loads((root / "PROVENANCE.json").read_text(encoding="utf-8"))
    assert provenance["commit"] == "b65272a18aef891941451705ff93e8a287f3eb72"
    assert provenance["version"] == OFFICECLI_VERSION
    assert provenance["binary_sha256"] == OFFICECLI_SHA256
    assert provenance["license"] == "Apache-2.0"
    assert {entry["path"] for entry in provenance["files"]} >= {
        "LICENSE", "NOTICE", "THIRD-PARTY-NOTICES.txt", "upstream/SKILL.md",
        "upstream/docx.md", "upstream/xlsx.md", "upstream/pptx.md",
    }
    for entry in provenance["files"]:
        assert hashlib.sha256((root / entry["path"]).read_bytes()).hexdigest() == entry["sha256"]
        assert provenance["commit"] in entry["source"]


LINUX = pytest.mark.skipif(sys.platform != "linux", reason="OfficeCLI 需在真实 Linux bubblewrap 中验收")


@LINUX
def test_real_officecli_reads_readonly_input_and_modifies_only_scratch(tmp_path):
    workspace = create_workspace(tmp_path, {})
    metadata = stage_officecli_skill(workspace)
    document = Document()
    document.add_paragraph("营业收入为100万元。")
    document.save(workspace / "input/main.docx")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "项目审查"
    presentation.save(workspace / "input/deck.pptx")
    originals = {name: hashlib.sha256((workspace / "input" / name).read_bytes()).hexdigest()
                 for name in ("main.docx", "deck.pptx")}
    sandbox = WorkspaceSandbox(Settings(_env_file=None), workspace)

    async def run():
        async def command(*arguments):
            result = await sandbox.execute("run_command", {"argv": ["officecli", *arguments]})
            assert result["exit_code"] == 0, result
            return result["stdout"]

        try:
            assert (await command("--version")).strip() == OFFICECLI_VERSION
            assert "营业收入为100万元。" in await command("view", "/workspace/input/main.docx", "text", "--json")
            assert "项目审查" in await command("view", "/workspace/input/deck.pptx", "text", "--json")
            blocked = await sandbox.execute("run_command", {"argv": ["officecli", "set",
                "/workspace/input/main.docx", "/body/p[1]", "--prop", "text=篡改"]})
            assert blocked["exit_code"] != 0, blocked
            for suffix in ("docx", "xlsx", "pptx"):
                await command("create", f"/workspace/scratch/work.{suffix}")
            await command("add", "/workspace/scratch/work.docx", "/body", "--type", "paragraph", "--prop", "text=临时核对")
            assert "临时核对" in await command("view", "/workspace/scratch/work.docx", "text", "--json")
            await command("set", "/workspace/scratch/work.xlsx", "/Sheet1/A1", "--prop", "value=42")
            assert '"42"' in await command("get", "/workspace/scratch/work.xlsx", "/Sheet1/A1", "--json")
            await command("add", "/workspace/scratch/work.pptx", "/", "--type", "slide", "--prop", "title=临时页面")
            assert "临时页面" in await command("view", "/workspace/scratch/work.pptx", "text", "--json")
            limits = await sandbox.execute("run_command", {"argv": ["python", "-c",
                "import resource; print(resource.getrlimit(resource.RLIMIT_AS)[1])"]})
            assert limits["stdout"].strip() == str(512 * 1024 * 1024)
            assert (await sandbox.execute("read_file", {"path": metadata["path"]}))["content"].startswith("---")
            for name, digest in originals.items():
                assert hashlib.sha256((workspace / "input" / name).read_bytes()).hexdigest() == digest
        finally:
            await sandbox.close()

    asyncio.run(run())
    assert not workspace.exists()


@LINUX
def test_real_officecli_computed_formula_values_never_become_saved_cache(tmp_path):
    workspace = create_workspace(tmp_path, {})
    path = workspace / "input/formula.xlsx"
    book = Workbook()
    sheet = book.active
    sheet["A1"], sheet["A2"], sheet["A3"], sheet["A4"] = 10, 20, "=SUM(A1:A2)", "=2+2"
    book.save(path)
    with zipfile.ZipFile(path) as source:
        entries = {info.filename: source.read(info.filename) for info in source.infolist()}
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(
        b"<f>SUM(A1:A2)</f><v></v>", b"<f>SUM(A1:A2)</f><v>999</v>"
    )
    with zipfile.ZipFile(path, "w") as target:
        for name, data in entries.items():
            target.writestr(name, data)
    original = path.read_bytes()
    sandbox = WorkspaceSandbox(Settings(_env_file=None), workspace)

    async def run():
        try:
            observed = []
            for cell in ("A3", "A4"):
                result = await sandbox.execute("run_command", {"argv": ["officecli", "get",
                    "/workspace/input/formula.xlsx", f"/Sheet/{cell}", "--json"]})
                assert result["exit_code"] == 0, result
                observed.append(json.loads(result["stdout"])["data"]["results"][0])
            assert observed[0]["format"]["cachedValue"] == "999"
            assert observed[0]["format"]["computedValue"] == "30"
            assert observed[1]["text"] == "4" and "cachedValue" not in observed[1]["format"]
            assert path.read_bytes() == original
            # The application parser continues to record missing cache rather than
            # promoting OfficeCLI's computed fallback into an original fact.
            from app.parsers import parse_document
            parsed = await parse_document(path, "formula", "formula.xlsx", Settings(_env_file=None), tmp_path / "parsed")
            assert any(warning.code == "formula_cache_missing" for warning in parsed.warnings)
            assert not any(block.text.strip() == "4" for block in parsed.blocks)
        finally:
            await sandbox.close()

    asyncio.run(run())

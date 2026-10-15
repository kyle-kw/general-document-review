"""Review artifacts, including evidence-anchored Word comments."""
from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.parsers import NS

CATEGORIES = {"typo": "错别字", "punctuation": "标点", "numeric": "数值一致性", "semantic": "语义一致性", "custom": "专项审查"}
SEVERITIES = {"high": "高", "medium": "中", "low": "低"}
FINDING_KINDS = {"violation": "内容不符", "missing": "内容缺失"}


def _label(location: dict[str, Any]) -> str:
    items = [("page", "页"), ("slide", "幻灯片"), ("sheet", "工作表"), ("cell", "单元格"),
             ("paragraph", "段落"), ("row", "行"), ("column", "列"), ("shape_name", "形状"),
             ("table", "表格"), ("part", "部件")]
    return "；".join(f"{label} {location[key]}" for key, label in items if key in location)


def _split_run(run: Run, offset: int) -> Run:
    """Split a plain run in-place without losing its formatting or XML attributes."""
    text = run.text
    following = deepcopy(run._r)
    run._r.addnext(following)
    trailing = Run(following, run._parent)
    run.text = text[:offset]
    trailing.text = text[offset:]
    return trailing


def _plain_run(run: Run) -> bool:
    for child in run._r:
        if child.tag not in {qn("w:rPr"), qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")}:
            return False
        if child.tag == qn("w:br") and child.get(qn("w:type")) not in (None, "textWrapping"):
            return False  # Splitting a page/column break would turn it into an ordinary line break.
    return True


def _anchor(document: Any, evidence: dict[str, Any], block: dict[str, Any]) -> tuple[list[Run], str]:
    if evidence.get("purpose", "quote") != "quote":
        return [], "定位上下文不能作为实际问题引文批注"
    location = block.get("location", {})
    if not location.get("commentable", False) or location.get("part") != "word/document.xml":
        return [], "此位置不支持 Word 原生批注（如页眉、页脚、图表或文本框）"
    xpath = location.get("xpath", "")
    nodes = document.element.getroottree().xpath(xpath, namespaces=NS) if xpath else []
    if len(nodes) != 1 or nodes[0].tag != qn("w:p"):
        return [], "无法重新定位原文段落"
    paragraph = Paragraph(nodes[0], document)
    quote = evidence.get("quote", "")
    if not quote:
        return [], "证据引文为空"
    start = evidence.get("start")
    if not isinstance(start, int) or paragraph.text[start:start + len(quote)] != quote:
        start = paragraph.text.find(quote)
        if start < 0:
            return [], "原文中无法找到证据引文"
        if paragraph.text.find(quote, start + 1) >= 0:
            return [], "引文在同一段落重复且未给出唯一字符位置"
    end = start + len(quote)
    # Complex inline structures are not rewritten. Use an intact direct run only when it contains the full quote.
    if not location.get("simple_runs", False):
        cursor = 0
        for run in paragraph.iter_inner_content():
            text = run.text
            if isinstance(run, Run) and cursor <= start and end <= cursor + len(text) and _plain_run(run):
                return [run], "使用完整原文 run 作为安全锚点"
            cursor += len(text)
        return [], "复杂内联结构没有安全批注锚点"
    cursor = 0
    for run in list(paragraph.runs):
        run_end = cursor + len(run.text)
        if run_end > start and cursor < end and not _plain_run(run):
            return [], "引文包含域、图片或其他无法安全拆分的 run"
        cursor = run_end
    # Cut at the end first so the original start offset stays valid.
    for boundary in (end, start):
        cursor = 0
        for run in paragraph.runs:
            length = len(run.text)
            if cursor < boundary < cursor + length:
                _split_run(run, boundary - cursor)
                break
            cursor += length
    runs: list[Run] = []
    cursor = 0
    for run in paragraph.runs:
        length = len(run.text)
        if length and cursor >= start and cursor + length <= end:
            runs.append(run)
        cursor += length
    return runs, "" if runs else "没有可批注的原文 run"


def _absence_labels(issue: dict[str, Any], documents: dict[str, dict[str, Any]]) -> list[str]:
    """Describe program-verified search coverage, never pretend absence is a quote."""
    labels = []
    missing_id = issue.get("missing_document_id")
    for coverage in issue.get("absence_coverage", []):
        document_id = coverage.get("document_id")
        filename = documents.get(document_id, {}).get("filename", "待核对文件")
        status = "已核对全部可提取内容" if coverage.get("complete") else "审查覆盖不完整，需人工确认"
        missing = "；未找到所需内容" if document_id == missing_id else ""
        labels.append(f"{filename}：{status}（{len(coverage.get('block_ids', []))} 个内容块，"
                      f"{len(coverage.get('reviewed_span_ids', []))} 个片段，"
                      f"{coverage.get('passes', 1)} 轮核查）{missing}")
    if missing_id and not any(coverage.get("document_id") == missing_id for coverage in issue.get("absence_coverage", [])):
        filename = documents.get(missing_id, {}).get("filename", "待核对文件")
        labels.append(f"{filename}：未找到所需内容；未提供完整覆盖信息，需人工确认")
    return labels


def _evidence_sides(issue: dict[str, Any], main_id: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    evidence = sorted(issue.get("evidence", []), key=lambda item: item.get("document_id") != main_id)
    missing_id = issue.get("missing_document_id") if issue.get("finding_kind") == "missing" else None
    if not missing_id:
        return evidence[:2], evidence[2:]
    # Give the absent document its own side. Its actual context stays explicit in
    # the extra evidence field rather than being misrepresented as the absence.
    present = [item for item in evidence if item.get("document_id") != missing_id]
    present.sort(key=lambda item: (item.get("document_id") != main_id, item.get("purpose") == "context"))
    sides = [*present[:1], {"document_id": missing_id, "absent": True}]
    sides.sort(key=lambda item: item.get("document_id") != main_id)
    extras = [item for item in evidence if item.get("document_id") == missing_id]
    return sides, [*present[1:], *extras]


def _annotate(report: dict[str, Any], main_path: Path, target: Path) -> None:
    document = Document(main_path)
    main_id = report["main_document_id"]
    documents = {doc["id"]: doc for doc in report.get("documents", [])}
    blocks = {block["id"]: block for doc in documents.values() for block in doc.get("blocks", [])}
    for issue in report.get("issues", []):
        issue["annotation_status"] = "unavailable"
        issue["annotation_reason"] = ("主文件缺少所需内容，没有真实主文件问题引文，无法批注缺失位置"
                                      if issue.get("finding_kind") == "missing" and issue.get("missing_document_id") == main_id
                                      else "没有真实主文件问题引文，无法批注")
        for evidence in issue.get("evidence", []):
            if evidence.get("document_id") != main_id or evidence.get("purpose", "quote") != "quote":
                continue
            block = blocks.get(evidence.get("block_id", ""))
            if not block:
                issue["annotation_reason"] = "证据内容块不存在"
                continue
            try:
                runs, reason = _anchor(document, evidence, block)
            except Exception as exc:
                issue["annotation_reason"] = f"无法安全定位批注（{type(exc).__name__}）"
                continue
            if not runs:
                issue["annotation_reason"] = reason
                continue
            details = [f"[{issue.get('id', '')}] {issue.get('title', '')}",
                       f"{CATEGORIES.get(issue.get('category'), '')} / {SEVERITIES.get(issue.get('severity'), '')}",
                       issue.get("description", ""), f"建议：{issue.get('suggestion', '')}"]
            if issue.get("category") == "custom":
                details.extend([f"审查要点：{issue.get('custom_rule_title', '')}",
                                f"问题性质：{FINDING_KINDS.get(issue.get('finding_kind', 'violation'), '')}"])
            for other in issue.get("evidence", []):
                source = documents.get(other.get("document_id"), {})
                source_block = blocks.get(other.get("block_id", ""), {})
                purpose = "定位上下文" if other.get("purpose") == "context" else "证据"
                details.append(f"{purpose}：{source.get('filename', other.get('document_id', ''))} "
                               f"{_label(source_block.get('location', other.get('location', {})))}\n{other.get('quote', '')}")
            if issue.get("finding_kind") == "missing":
                details.extend(f"缺失核对：{label}" for label in _absence_labels(issue, documents))
            try:
                document.add_comment(runs, text="\n".join(details), author="文件审查", initials="DR")
            except Exception as exc:
                issue["annotation_reason"] = f"原生批注添加失败（{type(exc).__name__}）"
                continue
            issue["annotation_status"] = "annotated"
            issue["annotation_reason"] = reason
            break
    temporary = target.with_suffix(".tmp.docx")
    document.save(temporary)
    temporary.replace(target)


def _chunks(value: Any) -> list[str]:
    text = ILLEGAL_CHARACTERS_RE.sub("", "" if value is None else str(value))
    return [text[index:index + 30000] for index in range(0, max(len(text), 1), 30000)]


def _row(sheet: Any, values: list[Any]) -> None:
    chunks = [_chunks(value) for value in values]
    for part in range(max(len(item) for item in chunks)):
        sheet.append([item[part] if part < len(item) else "" for item in chunks])
        for cell in sheet[sheet.max_row]:
            # All user/model strings remain literal text, including formulas and Excel injection prefixes.
            cell.data_type = "s"


def _workbook(report: dict[str, Any], target: Path) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "审查汇总"
    issues = report.get("issues", [])
    docs = {doc["id"]: doc for doc in report.get("documents", [])}
    blocks = {block["id"]: block for doc in docs.values() for block in doc.get("blocks", [])}
    all_warnings = [(doc, warning) for doc in docs.values() for warning in doc.get("warnings", [])]
    incomplete = any(warning.get("incomplete", True) for _, warning in all_warnings) or any(
        issue.get("finding_kind") == "missing" and (
            not any(item.get("document_id") == issue.get("missing_document_id") for item in issue.get("absence_coverage", []))
            or any(not item.get("complete") for item in issue.get("absence_coverage", []))
        ) for issue in issues
    )
    summary_rows = [
        ["项目", "内容"], ["任务 ID", report.get("review_id", "")], ["状态", report.get("status", "")],
        ["主文件", docs.get(report.get("main_document_id"), {}).get("filename", "")],
        ["文件数", len(docs)], ["问题总数", len(issues)],
        ["待确认问题", sum(issue.get("confidence") == "needs_confirmation" for issue in issues)],
        ["审查覆盖", "不完整，详见解析与覆盖情况" if incomplete else "已完成可提取内容的约定审查"],
        ["审查范围", "主文件：错字、标点及内部数值/语义一致性；其他文件：仅与主文件的一致性；按整理后的专项要点核对相关内容；不做外部事实核查"],
        ["创建时间", report.get("created_at", "")], ["完成时间", report.get("finished_at", "")],
        ["长字段说明", "超过 Excel 单元格容量的字段按顺序拆分到后续行，JSON 保留完整结构。"],
    ]
    for values in summary_rows:
        _row(summary, values)
    requirements = report.get("review_requirements") or {}
    _row(summary, ["专项审查要点数", len(requirements.get("custom_points", []))])
    for index, rule in enumerate(requirements.get("custom_points", []), 1):
        _row(summary, [f"审查要点 {index}：{rule.get('title', '')}", rule.get("criterion", "")])
    for index, note in enumerate(requirements.get("notes", []), 1):
        _row(summary, [f"审查备注 {index}", note])
    if report.get("requirements_policy_version"):
        _row(summary, ["审查要求策略版本", report["requirements_policy_version"]])
    detail = workbook.create_sheet("问题明细")
    _row(detail, ["问题 ID", "类型", "范围", "严重程度", "确认状态", "标题", "问题说明", "建议", "主文件/证据一",
                  "位置一", "引文一", "关联文件/证据二", "位置二", "引文二", "其他证据", "批注状态", "批注说明",
                  "规则 ID", "审查要点", "问题性质", "缺失文件", "缺失内容核对覆盖"])
    for issue in issues:
        evidence, extra_evidence = _evidence_sides(issue, report.get("main_document_id", ""))
        cells: list[Any] = [issue.get("id"), CATEGORIES.get(issue.get("category"), issue.get("category")),
                            "跨文件" if issue.get("scope") == "cross" else "主文件内部",
                            SEVERITIES.get(issue.get("severity"), issue.get("severity")),
                            "待确认" if issue.get("confidence") == "needs_confirmation" else "明确问题",
                            issue.get("title"), issue.get("description"), issue.get("suggestion")]
        for index in range(2):
            item = evidence[index] if len(evidence) > index else {}
            block = blocks.get(item.get("block_id", ""), {})
            cells.extend([docs.get(item.get("document_id"), {}).get("filename", ""),
                          "" if item.get("absent") else _label(block.get("location", item.get("location", {}))),
                          "" if item.get("absent") else item.get("quote", "")])
        cells.extend([json.dumps(extra_evidence, ensure_ascii=False) if extra_evidence else "",
                      issue.get("annotation_status", "not_applicable"), issue.get("annotation_reason", ""),
                      issue.get("custom_rule_id", ""), issue.get("custom_rule_title", ""),
                      FINDING_KINDS.get(issue.get("finding_kind", "violation"), ""),
                      docs.get(issue.get("missing_document_id"), {}).get("filename", ""),
                      "\n".join(_absence_labels(issue, docs))])
        _row(detail, cells)
    coverage = workbook.create_sheet("解析与覆盖情况")
    _row(coverage, ["文件", "格式", "内容块数", "状态", "代码", "位置", "说明"])
    for doc in docs.values():
        if not doc.get("warnings"):
            _row(coverage, [doc["filename"], doc["format"], len(doc.get("blocks", [])), "完成", "", "", "无解析警告"])
        for warning in doc.get("warnings", []):
            _row(coverage, [doc["filename"], doc["format"], len(doc.get("blocks", [])),
                            "未完整覆盖" if warning.get("incomplete", True) else "范围说明", warning.get("code"),
                            _label(warning.get("location", {})), warning.get("message")])
    for issue in issues:
        if issue.get("finding_kind") != "missing":
            continue
        for item in issue.get("absence_coverage", []):
            doc = docs.get(item.get("document_id"), {})
            _row(coverage, [doc.get("filename", "待核对文件"), doc.get("format", ""), len(item.get("block_ids", [])),
                           "缺失核对完整" if item.get("complete") else "缺失核对不完整", "custom_absence_coverage", "",
                           f"{issue.get('custom_rule_title', '')}：已核对 {len(item.get('reviewed_span_ids', []))} 个片段，"
                           f"{item.get('passes', 1)} 轮核查；"
                           + ("未找到所需内容" if item.get("document_id") == issue.get("missing_document_id") else "对照内容核对覆盖")])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 26
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173D38")
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 26 if index < 6 else 45
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    summary.column_dimensions["B"].width = 100
    temporary = target.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(target)


def export_report(report: dict[str, Any], main_path: Path, output_dir: Path) -> list[dict[str, str]]:
    output_dir.mkdir(parents=True, exist_ok=True)
    artifacts: list[dict[str, str]] = []
    main = next((doc for doc in report.get("documents", []) if doc["id"] == report["main_document_id"]), {})
    if main.get("format", main_path.suffix.lstrip(".")).lower() == "docx":
        target = output_dir / "annotated-main.docx"
        _annotate(report, main_path, target)
        artifacts.append({"name": target.name, "path": str(target), "media_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "kind": "annotated"})
    else:
        for issue in report.get("issues", []):
            issue["annotation_status"] = "not_applicable"
            issue["annotation_reason"] = "主文件为非 DOCX 格式"
    json_path = output_dir / "report.json"
    temporary = json_path.with_suffix(".tmp.json")
    temporary.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    temporary.replace(json_path)
    excel_path = output_dir / "report.xlsx"
    _workbook(report, excel_path)
    return [
        {"name": json_path.name, "path": str(json_path), "media_type": "application/json", "kind": "json"},
        {"name": excel_path.name, "path": str(excel_path), "media_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "kind": "excel"},
        *artifacts,
    ]

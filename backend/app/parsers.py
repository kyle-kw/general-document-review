"""Deterministic source extraction. No parser executes document code or formulas."""
from __future__ import annotations

import asyncio
import base64
import hashlib
import io
import json
import posixpath
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

import httpx
import pdfplumber
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader, PdfWriter

from app.config import Settings
from app.models import Block, CoverageWarning, ParsedDocument

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


class FileParsingError(ValueError):
    """A source cannot be safely opened; retrying the same source will not help."""


class ParserCancelled(RuntimeError):
    """Cooperative cancellation between parser work units."""


def _check_cancelled(cancelled: Callable[[], bool] | None) -> None:
    if cancelled is not None and cancelled():
        raise ParserCancelled("审查已取消")


def _xml(data: bytes) -> etree._Element:
    return etree.fromstring(data, parser=etree.XMLParser(resolve_entities=False, no_network=True))


def _add(doc: ParsedDocument, text: str, kind: str, location: dict[str, Any]) -> None:
    if not text.strip():
        return
    identity = json.dumps(location, ensure_ascii=False, sort_keys=True, default=str)
    digest = hashlib.sha256(f"{doc.id}:{identity}".encode()).hexdigest()[:20]
    doc.blocks.append(Block(id=f"{doc.id}-{digest}", document_id=doc.id,
                            text=text, kind=kind, location=location))


def _warning(doc: ParsedDocument, code: str, message: str,
             location: dict[str, Any] | None = None, incomplete: bool = True) -> None:
    doc.warnings.append(CoverageWarning(code=code, message=message,
                                       location=location or {}, incomplete=incomplete))


def _check_archive(path: Path, settings: Settings) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            total = sum(item.file_size for item in archive.infolist())
            if total > settings.max_unpacked_mb * 1024 * 1024:
                raise FileParsingError("Office 解压大小超过配置限制")
            if any(item.flag_bits & 1 for item in archive.infolist()):
                raise FileParsingError("不支持密码保护的 Office 文件")
    except zipfile.BadZipFile as exc:
        raise FileParsingError("Office 文件损坏或已加密") from exc


def _charts(path: Path, doc: ParsedDocument, cancelled: Callable[[], bool] | None = None) -> None:
    """Read native chart caches, including charts in Word and Excel."""
    with zipfile.ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            _check_cancelled(cancelled)
            if not re.search(r"/charts/chart\d+\.xml$", name):
                continue
            try:
                root = _xml(archive.read(name))
                title = "".join(root.xpath("./c:chart/c:title//a:t/text()", namespaces=NS))
                references = [block.location for block in doc.blocks
                              if block.kind == "chart_reference" and block.location.get("part") == name]
                source_location = {key: value for key, value in references[0].items()
                                   if key not in {"part", "commentable"}} if references else {}
                series = root.xpath("//c:ser", namespaces=NS)
                for index, series_node in enumerate(series):
                    lines: list[str] = []
                    if title:
                        lines.append(f"图表标题：{title}")
                    for child in series_node:
                        tag = etree.QName(child).localname
                        if tag in {"idx", "order", "spPr", "dPt", "extLst"}:
                            continue
                        points = child.xpath(".//c:pt", namespaces=NS)
                        values = ["".join(p.xpath("./c:v/text()", namespaces=NS)) for p in points]
                        if not values:
                            values = child.xpath("./c:v/text()", namespaces=NS)
                        if values:
                            lines.append(f"{tag}：" + " | ".join(values))
                        refs = child.xpath(".//c:f/text()", namespaces=NS)
                        if refs and not values:
                            _warning(doc, "chart_cache_missing", "图表没有保存数据缓存，无法审查引用数据",
                                     {"part": name, "series": index + 1, "references": refs})
                    _add(doc, "\n".join(lines), "chart",
                         {**source_location, "part": name, "chart_title": title,
                          "series": index + 1, "commentable": False})
            except (etree.XMLSyntaxError, ValueError) as exc:
                _warning(doc, "chart_parse_failed", f"图表解析失败：{exc}", {"part": name})


def _scope_warning(path: Path, doc: ParsedDocument) -> None:
    with zipfile.ZipFile(path) as archive:
        media = [name for name in archive.namelist() if "/media/" in name]
        if media:
            _warning(doc, "embedded_images_excluded", f"{len(media)} 个内嵌媒体文件未进行 OCR（审查范围约定）",
                     incomplete=False)
        chart_workbooks: set[str] = set()
        for name in archive.namelist():
            if re.search(r"/charts/_rels/chart\d+\.xml\.rels$", name):
                root = _xml(archive.read(name))
                for relationship in root:
                    if relationship.get("TargetMode") == "External":
                        continue
                    target = relationship.get("Target", "")
                    # Relationship targets are relative to the corresponding chart part.
                    chart_workbooks.add(posixpath.normpath(posixpath.join(posixpath.dirname(posixpath.dirname(name)), target)))
        if any("/embeddings/" in name and name not in chart_workbooks for name in archive.namelist()):
            _warning(doc, "embedded_objects_excluded", "嵌入对象未执行或打开；原生图表只读取保存的数据缓存")
        if any("/diagrams/" in name and name.endswith(".xml") for name in archive.namelist()):
            _warning(doc, "smartart_excluded", "SmartArt 数据和布局未展开审查；仅提取其可直接读取的文本")


def _docx(path: Path, doc: ParsedDocument, cancelled: Callable[[], bool] | None = None) -> None:
    document = Document(path)
    with zipfile.ZipFile(path) as archive:
        parts = ["word/document.xml"] + sorted(
            name for name in archive.namelist()
            if re.match(r"word/(?:header\d+|footer\d+|footnotes|endnotes)\.xml$", name)
        )
        for part in parts:
            _check_cancelled(cancelled)
            root = document.element if part == "word/document.xml" else parse_xml(archive.read(part))
            tree = root.getroottree()
            for number, paragraph_node in enumerate(root.iter(qn("w:p")), 1):
                _check_cancelled(cancelled)
                # Paragraph.text ignores nested text boxes; their own paragraphs are visited separately.
                paragraph = Paragraph(paragraph_node, document)
                text = paragraph.text
                tracked = paragraph_node.xpath(".//w:ins | .//w:del") if part == "word/document.xml" else []
                if tracked:
                    _warning(doc, "tracked_changes", "段落包含修订，提取内容可能不涵盖所有修订状态",
                             {"part": part, "paragraph": number})
                if not text.strip():
                    continue
                ancestors = {etree.QName(node).localname for node in paragraph_node.iterancestors()}
                runs, cursor = [], 0
                for run in paragraph.runs:
                    runs.append({"start": cursor, "end": cursor + len(run.text),
                                 "xpath": tree.getpath(run._r)})
                    cursor += len(run.text)
                simple = cursor == len(text) and all(
                    child.tag in {qn("w:r"), qn("w:pPr"), qn("w:bookmarkStart"), qn("w:bookmarkEnd"),
                                  qn("w:commentRangeStart"), qn("w:commentRangeEnd")}
                    for child in paragraph_node
                )
                location = {
                    "part": part, "xpath": tree.getpath(paragraph_node), "paragraph": number,
                    "char_start": 0, "char_end": len(text), "runs": runs,
                    "commentable": part == "word/document.xml" and "txbxContent" not in ancestors,
                    "simple_runs": simple,
                }
                if "tc" in ancestors:
                    cell_node = next(node for node in paragraph_node.iterancestors() if node.tag == qn("w:tc"))
                    row_node = cell_node.getparent()
                    table_node = row_node.getparent()
                    rows = [node for node in table_node if node.tag == qn("w:tr")]
                    cells = [node for node in row_node if node.tag == qn("w:tc")]
                    location.update({"table_cell": tree.getpath(cell_node), "row": rows.index(row_node) + 1,
                                     "column": cells.index(cell_node) + 1})
                    if rows:
                        header_cells = [node for node in rows[0] if node.tag == qn("w:tc")]
                        column = cells.index(cell_node)
                        if column < len(header_cells):
                            location["column_header"] = " ".join(header_cells[column].xpath(".//w:t/text()"))
                    if cells:
                        location["row_header"] = " ".join(cells[0].xpath(".//w:t/text()"))
                if "txbxContent" in ancestors:
                    location["commentable"] = False
                _add(doc, text, "table_cell" if "tc" in ancestors else "paragraph", location)
    _charts(path, doc, cancelled)
    _scope_warning(path, doc)


def _value(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _xlsx(path: Path, doc: ParsedDocument, cancelled: Callable[[], bool] | None = None) -> None:
    formulas = load_workbook(path, data_only=False, keep_links=False)
    cached = load_workbook(path, data_only=True, keep_links=False)
    try:
        for sheet in formulas:
            _check_cancelled(cancelled)
            values = cached[sheet.title]
            for cell in sorted(sheet._cells.values(), key=lambda c: (c.row, c.column)):
                _check_cancelled(cancelled)
                if cell.value is None:
                    continue
                formula = cell.data_type == "f"
                value = values[cell.coordinate].value if formula else cell.value
                location: dict[str, Any] = {
                    "sheet": sheet.title, "cell": cell.coordinate, "row": cell.row,
                    "column": cell.column, "number_format": cell.number_format,
                    "sheet_state": sheet.sheet_state, "commentable": False,
                    "column_header": _value(sheet.cell(1, cell.column).value or ""),
                    "row_header": _value(sheet.cell(cell.row, 1).value or ""),
                }
                for region in sheet.merged_cells.ranges:
                    if cell.coordinate in region:
                        location["merged_range"] = str(region)
                        break
                if formula:
                    location["formula"] = cell.value
                    if value is None:
                        _warning(doc, "formula_cache_missing", "公式没有保存计算值，未将其作为零或计算结果审查", location)
                        continue
                _add(doc, _value(value), "cell", location)
            if sheet._images:
                _warning(doc, "sheet_images_excluded", "工作表中的图片未进行 OCR", {"sheet": sheet.title}, False)
    finally:
        formulas.close()
        cached.close()
    _charts(path, doc, cancelled)
    _scope_warning(path, doc)


def _pptx(path: Path, doc: ParsedDocument, cancelled: Callable[[], bool] | None = None) -> None:
    presentation = Presentation(path)

    def shapes(items: Any, slide_number: int, parents: list[int]) -> None:
        for shape in items:
            _check_cancelled(cancelled)
            shape_path = parents + [shape.shape_id]
            location = {"slide": slide_number, "shape": shape.shape_id, "shape_path": shape_path,
                        "shape_name": shape.name, "commentable": False}
            if hasattr(shape, "shapes"):
                shapes(shape.shapes, slide_number, shape_path)
            if shape.has_text_frame:
                for number, paragraph in enumerate(shape.text_frame.paragraphs, 1):
                    _add(doc, paragraph.text, "paragraph", {**location, "paragraph": number})
            if shape.has_table:
                for row_number, row in enumerate(shape.table.rows, 1):
                    for column_number, cell in enumerate(row.cells, 1):
                        if cell.is_spanned:
                            continue
                        _add(doc, cell.text, "table_cell", {**location, "row": row_number,
                                                           "column": column_number})
            if shape.has_chart:
                # Chart XML also keeps unsupported chart types and original cache values.
                chart_part = str(shape.chart.part.partname).lstrip("/")
                _add(doc, f"原生图表：{chart_part}", "chart_reference", {**location, "part": chart_part})

    for slide_number, slide in enumerate(presentation.slides, 1):
        _check_cancelled(cancelled)
        shapes(slide.shapes, slide_number, [])
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                for number, paragraph in enumerate(frame.paragraphs, 1):
                    _add(doc, paragraph.text, "notes", {"slide": slide_number, "notes_paragraph": number,
                                                        "commentable": False})
    _charts(path, doc, cancelled)
    _scope_warning(path, doc)


def _pdf_reader(path: Path) -> PdfReader:
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise FileParsingError("不支持密码保护的 PDF 文件")
    return reader


def _pdf_local(path: Path, doc: ParsedDocument, cancelled: Callable[[], bool] | None = None) -> None:
    _pdf_reader(path)
    with pdfplumber.open(path) as pdf:
        for number, page in enumerate(pdf.pages, 1):
            _check_cancelled(cancelled)
            try:
                tables = page.find_tables()
                for table_index, table in enumerate(tables, 1):
                    for row_index, row in enumerate(table.extract(), 1):
                        _add(doc, " | ".join(value or "" for value in row), "table_row",
                             {"page": number, "table": table_index, "row": row_index,
                              "bbox": list(table.bbox), "commentable": False})
                for index, line in enumerate(page.extract_text_lines(), 1):
                    x, y = (line["x0"] + line["x1"]) / 2, (line["top"] + line["bottom"]) / 2
                    if any(t.bbox[0] <= x <= t.bbox[2] and t.bbox[1] <= y <= t.bbox[3] for t in tables):
                        continue
                    _add(doc, line["text"], "paragraph", {"page": number, "line": index,
                         "bbox": [line["x0"], line["top"], line["x1"], line["bottom"]], "commentable": False})
                if not page.chars:
                    _warning(doc, "pdf_no_text", "页面没有可提取文字，可能为空白页或扫描页；未配置 OCR",
                             {"page": number})
                if any("\ufffd" in b.text or "(cid:" in b.text for b in doc.blocks if b.location.get("page") == number):
                    _warning(doc, "pdf_text_encoding", "页面文字存在缺失字符或字体编码异常", {"page": number})
            except Exception as exc:
                _warning(doc, "pdf_page_failed", f"页面提取失败：{type(exc).__name__}", {"page": number})


async def _retry_delay(delay: float, cancelled: Callable[[], bool] | None) -> None:
    while delay > 0:
        _check_cancelled(cancelled)
        interval = min(delay, 0.2)
        await asyncio.sleep(interval)
        delay -= interval


async def _pdf_ocr(path: Path, doc: ParsedDocument, settings: Settings, work_dir: Path,
                   cancelled: Callable[[], bool] | None = None) -> None:
    reader = _pdf_reader(path)
    work_dir.mkdir(parents=True, exist_ok=True)
    endpoint = settings.paddleocr_url.rstrip("/")
    if not endpoint.endswith("/layout-parsing"):
        endpoint += "/layout-parsing"
    headers = {"Authorization": f"token {settings.paddleocr_token.get_secret_value()}"} if settings.paddleocr_token.get_secret_value() else {}
    async with httpx.AsyncClient(timeout=settings.paddleocr_timeout, follow_redirects=False) as client:
        for offset in range(0, len(reader.pages), settings.paddleocr_batch_pages):
            _check_cancelled(cancelled)
            count = min(settings.paddleocr_batch_pages, len(reader.pages) - offset)
            writer, buffer = PdfWriter(), io.BytesIO()
            for index in range(offset, offset + count):
                _check_cancelled(cancelled)
                writer.add_page(reader.pages[index])
            writer.write(buffer)
            payload = {"file": base64.b64encode(buffer.getvalue()).decode("ascii"), "fileType": 0,
                       "visualize": False, "restructurePages": False}
            raw_path = work_dir / f"ocr-{doc.id}-{offset + 1:06d}.json"
            try:
                for attempt in range(settings.llm_max_retries + 1):
                    _check_cancelled(cancelled)
                    try:
                        response = await client.post(endpoint, headers=headers, json=payload)
                    except httpx.RequestError:
                        if attempt == settings.llm_max_retries:
                            raise
                        await _retry_delay(2 ** attempt, cancelled)
                        continue
                    _check_cancelled(cancelled)
                    # Preserve server responses even on errors. Never follow returned URLs or fetch media.
                    raw_path.write_text(response.text, encoding="utf-8")
                    if response.status_code == 429 or 500 <= response.status_code < 600:
                        if attempt < settings.llm_max_retries:
                            raw_path.with_name(f"{raw_path.stem}-attempt{attempt + 1}.json").write_text(response.text, encoding="utf-8")
                            await _retry_delay(2 ** attempt, cancelled)
                            continue
                    break
                response.raise_for_status()
                data = response.json()
                if data.get("errorCode", 0) not in (0, None):
                    raise ValueError("OCR 服务返回错误码")
                pages = data["result"]["layoutParsingResults"]
                if not isinstance(pages, list):
                    raise ValueError("OCR 响应缺少页面数组")
                if len(pages) != count:
                    _warning(doc, "ocr_page_count_mismatch", f"OCR 返回 {len(pages)} 页，预期 {count} 页",
                             {"page_start": offset + 1, "page_end": offset + count, "raw_response": raw_path.name})
                # Protocol results are ordered and omit page_index. Excess results are never assigned a fake page.
                for relative, page in enumerate(pages[:count]):
                    _check_cancelled(cancelled)
                    original_page = offset + relative + 1
                    explicit_index = page.get("page_index", page.get("prunedResult", {}).get("page_index"))
                    if explicit_index is not None and explicit_index != relative:
                        _warning(doc, "ocr_page_index_mismatch", "OCR 页码索引不符合批次顺序，未使用该页",
                                 {"page": original_page})
                        continue
                    text = page.get("markdown", {}).get("text", "")
                    if not isinstance(text, str) or not text.strip():
                        _warning(doc, "ocr_empty_page", "OCR 页面没有返回文字", {"page": original_page})
                        continue
                    for paragraph_index, paragraph in enumerate(re.split(r"\n\s*\n", text), 1):
                        _add(doc, paragraph, "ocr_markdown", {"page": original_page, "paragraph": paragraph_index,
                             "raw_response": raw_path.name, "commentable": False})
                for missing in range(len(pages), count):
                    _warning(doc, "ocr_missing_page", "OCR 未返回该页，无法完成审查", {"page": offset + missing + 1})
            except (httpx.HTTPError, ValueError, KeyError, TypeError) as exc:
                _warning(doc, "ocr_batch_failed", f"OCR 批次失败：{type(exc).__name__}；未切换本地解析",
                         {"page_start": offset + 1, "page_end": offset + count, "raw_response": raw_path.name})


async def parse_document(path: Path, document_id: str, filename: str,
                         settings: Settings, work_dir: Path,
                         cancelled: Callable[[], bool] | None = None) -> ParsedDocument:
    extension = Path(filename).suffix.lower().lstrip(".")
    if extension not in {"docx", "xlsx", "pptx", "pdf"}:
        raise FileParsingError("仅支持 docx、xlsx、pptx 和 pdf 格式")
    doc = ParsedDocument(id=document_id, filename=filename, format=extension)
    try:
        _check_cancelled(cancelled)
        if extension != "pdf":
            await asyncio.to_thread(_check_archive, path, settings)
            await asyncio.to_thread({"docx": _docx, "xlsx": _xlsx, "pptx": _pptx}[extension], path, doc, cancelled)
        elif settings.paddleocr_url:
            await _pdf_ocr(path, doc, settings, work_dir, cancelled)
        else:
            await asyncio.to_thread(_pdf_local, path, doc, cancelled)
        _check_cancelled(cancelled)
    except (FileParsingError, ParserCancelled):
        raise
    except Exception as exc:
        raise FileParsingError(f"无法解析 {filename}（{type(exc).__name__}）") from exc
    if not doc.blocks:
        _warning(doc, "no_extractable_content", "文件没有可审查内容")
    return doc
